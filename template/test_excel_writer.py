# Databricks notebook source
# MAGIC %md
# MAGIC ###Installing dependencies

# COMMAND ----------

# MAGIC %pip install openpyxl

# COMMAND ----------

# MAGIC %md
# MAGIC ### Importing libraries

# COMMAND ----------

import pandas as pd
import os
from datetime import datetime

import tempfile

import openpyxl
from openpyxl.styles import Font, Border, Side, PatternFill, Alignment
from openpyxl.utils import get_column_letter

# COMMAND ----------

# MAGIC %md
# MAGIC ###Output file path

# COMMAND ----------

mnt = "/mnt/projects/FinancialSpreading"

output_file_path = "dbfs:" + mnt + "/Test/output_test_file.xlsx"

# COMMAND ----------

# MAGIC %md
# MAGIC ###Excel writing function

# COMMAND ----------

def excel_write(dataframes:dict[str, pd.DataFrame], output_path:str, overwrite: bool= False):
    
    # File naming convention
    if output_path is None:
        raise ValueError("Output path cannot be None")
    if not output_path.endswith('.xlsx'):
        raise ValueError("Output path must be an Excel file")
    if not dataframes:
        raise ValueError("Sheet configurations cannot be empty")

    if overwrite:
        if os.path.exists(output_path):
            os.remove(output_path)

    file_name = os.path.basename(output_path)
    

    with tempfile.TemporaryDirectory(prefix="fin_spreading_") as temp_dir:
        temp_file_path = os.path.join(temp_dir, file_name)
        with pd.ExcelWriter(temp_file_path, engine='openpyxl') as writer:
            for sheet_name, data in dataframes.items():
                data.to_excel(writer, sheet_name=sheet_name, index=False)
                worksheet = writer.sheets[sheet_name]

                wrap_columns = set()
                for idx, col in enumerate(data.columns, 1):
                    max_length = max(
                        data[col].astype(str).map(len).max(),
                        len(str(col))
                    )
                    adjusted_length = max_length + 3
                    worksheet.column_dimensions[get_column_letter(idx)].width = min(adjusted_length, 50)
                    if adjusted_length > 50:
                        wrap_columns.add(idx)
                
                worksheet.freeze_panes = "B2"
                
                for cell in worksheet[1]:
                    cell.font = Font(name="Arial", size=12, bold=True, color="FFFFFF")
                    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=cell.column in wrap_columns)
                    cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")

                thin_side = Side(style="thin", color="000000")
                thin_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

                for row in worksheet.iter_rows():
                    for cell in row:
                        cell.border = thin_border
                        if cell.column in wrap_columns and cell.row > 1:
                            cell.alignment = Alignment(wrap_text=True)
        
        if os.path.exists(temp_file_path):
            dbutils.fs.cp(f"file://{temp_file_path}", output_path)
    
    print(f"Excel file created: {output_path}")

# COMMAND ----------

# MAGIC %md
# MAGIC ### Test function

# COMMAND ----------

# income_data = pd.DataFrame(
#     [
#         {"Variable": "Net Sales", "Value": 3500000, "Score": 0.96},
#         {"Variable": "COGS", "Value": 2100000, "Score": 0.94},
#         {"Variable": "Gross Profit", "Value": 1400000, "Score": 0.95},
#         {"Variable": "Operating Expenses                                                  3", "Value": 850000, "Score": 0.89},
#         {"Variable": "Net Income                                dfgfdhhhhhhhhhhhhhhhhhhhhhhhhhhhhhhhhhhhhhhhhhhhhhhhh                                                     1", "Value": 550000, "Score": 0.92}
#     ]
# )
    
# balance_data = pd.DataFrame(
#     [
#         {"Variable": "Total Assets", "Value": 8500000, "Score": 0.97},
#         {"Variable": "Liabilities", "Value": 3200000, "Score": 0.93},
#         {"Variable": "Net Worth", "Value": 5300000, "Score": 0.91},
#         {"Variable": "Cash", "Value": 1200000, "Score": 0.98}
#     ]
# )

# sheet_configs = {
#     "Sheet1": pd.DataFrame({"Column1": [1, 2, 3], "Column2": ["A", "B", "C"]}),
#     "Sheet2": pd.DataFrame({"Column1": [4, 5, 6], "Column2": ["D", "E", "F"]}),
#     "Income": income_data,
#     "Balance": balance_data
# }

# COMMAND ----------

# excel_write(dataframes = sheet_configs, output_path = output_file_path, overwrite = True)
