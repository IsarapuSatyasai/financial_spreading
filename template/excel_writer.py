"""Defining function that takes dataframe dictionary as input and writes the output to the excel file."""

import pandas as pd
import os

# from databricks.sd    k.runtime import dbutils
import tempfile

import openpyxl
from openpyxl.styles import Font, Border, Side, PatternFill, Alignment
from openpyxl.utils import get_column_letter

import shutil

if os.path.exists(temp_file_path):
    shutil.copy2(
        temp_file_path,
        output_path
    )

def excel_write(dataframes:dict[str, pd.DataFrame], output_path:str, overwrite: bool= False) -> str:
    """
    This function takes dictionary of pandas dataframes and writes results to the output path.
    """
    
    # File naming convention
    if output_path is None:
        raise ValueError("Output path cannot be None")
    if not output_path.endswith('.xlsx'):
        raise ValueError("Output path must be an Excel file")
    if not dataframes:
        raise ValueError("Sheet configurations cannot be empty")

    if overwrite:
        if os.path.exists(output_path):
            os.remove(output_path)

    file_name = os.path.basename(output_path)
    

    with tempfile.TemporaryDirectory(prefix="fin_spreading_") as temp_dir:
        temp_file_path = os.path.join(temp_dir, file_name)
        with pd.ExcelWriter(temp_file_path, engine='openpyxl') as writer:
            for sheet_name, data in dataframes.items():
                data.to_excel(writer, sheet_name=sheet_name, index=False)
                worksheet = writer.sheets[sheet_name]

                wrap_columns = set()
                for idx, col in enumerate(data.columns, 1):
                    max_length = max(
                        data[col].astype(str).map(len).max(),
                        len(str(col))
                    )
                    adjusted_length = max_length + 3
                    worksheet.column_dimensions[get_column_letter(idx)].width = min(adjusted_length, 50)
                    if adjusted_length > 50:
                        wrap_columns.add(idx)
                
                worksheet.freeze_panes = "B2"
                
                for cell in worksheet[1]:
                    cell.font = Font(name="Arial", size=12, bold=True, color="FFFFFF")
                    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=cell.column in wrap_columns)
                    cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")

                thin_side = Side(style="thin", color="000000")
                thin_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

                for row in worksheet.iter_rows():
                    for cell in row:
                        cell.border = thin_border
                        if cell.column in wrap_columns and cell.row > 1:
                            cell.alignment = Alignment(wrap_text=True)
        
        if os.path.exists(temp_file_path):
            dbutils.fs.cp(f"file://{temp_file_path}", output_path)
    
    return f"Excel file created: {output_path}"